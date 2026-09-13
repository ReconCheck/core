"""Tabular parsing (CSV / TSV / XLSX).

The parser's job is to turn a (messy) table into a :class:`Table` where every
cell carries a :class:`SourceLoc`. Known limitations (see DESIGN.md): merged
cells in XLSX yield the value from the top-left cell only; PDF and image
parsing is a later milestone.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..errors import UnsupportedFormatError
from ..models import Cell, Document, Row, SourceLoc, Table

_DELIMITERS = {".csv": ",", ".tsv": "\t", ".txt": None}


def load_document(path: str | Path, sheet: str | None = None) -> Document:
    """Parse ``path`` into a :class:`Document`.

    ``sheet`` restricts XLSX input to one worksheet by title.
    """
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        return _document_from_delimited(p)
    if suffix in {".xlsx", ".xlsm"}:
        return _document_from_xlsx(p, sheet)
    raise UnsupportedFormatError(
        f"unsupported file type '{suffix}' (supported: .csv, .tsv, .txt, .xlsx)"
    )


def _decode(raw: bytes) -> str:
    """Decode bytes, trying encodings Chinese Excel exports actually use."""
    for encoding in ("utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _sniff_delimiter(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text[:2048], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _dedupe_headers(headers: list[str]) -> list[str]:
    """Ensure header names are unique; blank ones get a generated name."""
    seen: dict[str, int] = {}
    out: list[str] = []
    for header in headers:
        if not header:
            header = f"col{len(out) + 1}"
        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 0
        out.append(header)
    return out


def _row_from_values(values: list[str], path: Path, row_num: int, headers: list[str]) -> Row:
    cells: list[Cell] = []
    for j, value in enumerate(values):
        col = j + 1
        header = headers[j] if j < len(headers) else f"col{col}"
        cell = Cell(
            text=str(value).strip(),
            loc=SourceLoc(path=str(path), row=row_num, col=col, header=header),
        )
        cell.coerce()
        cells.append(cell)
    return Row(cells)


def _document_from_delimited(p: Path) -> Document:
    text = _decode(p.read_bytes())
    delimiter = _DELIMITERS[p.suffix.lower()] or _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [["" if value is None else value for value in row] for row in reader]
    if not rows:
        return Document(str(p), [])
    headers = _dedupe_headers([str(c).strip() for c in rows[0]])
    data_rows = rows[1:]
    width = max((len(row) for row in data_rows), default=len(headers))
    while len(headers) < width:
        headers.append(f"col{len(headers) + 1}")
    entries: list[Row] = []
    for i, raw in enumerate(data_rows, start=2):
        if not any(str(c).strip() for c in raw):
            continue
        entries.append(_row_from_values(raw, p, i, headers))
    table = Table(headers=headers, rows=entries, loc=SourceLoc(path=str(p), row=2, col=1))
    return Document(str(p), [table])


def _document_from_xlsx(p: Path, sheet: str | None = None) -> Document:
    wb = load_workbook(p, read_only=False, data_only=True)
    tables: list[Table] = []
    try:
        for ws in wb.worksheets:
            if sheet is not None and ws.title != sheet:
                continue
            grid: list[tuple[int, list[str]]] = []
            for i, values in enumerate(ws.iter_rows(values_only=True), start=ws.min_row):
                vals = ["" if v is None else str(v) for v in values]
                if not any(vals):
                    continue
                grid.append((i, vals))
            if not grid:
                continue
            header_row, raw_header = grid[0]
            headers = _dedupe_headers([str(h).strip() for h in raw_header])
            width = max((len(vals) for _, vals in grid), default=len(headers))
            while len(headers) < width:
                headers.append(f"col{len(headers) + 1}")
            entries = [
                _row_xlsx(vals, p, ws.title, row_num, headers) for row_num, vals in grid[1:]
            ]
            loc = SourceLoc(path=str(p), sheet=ws.title, row=header_row + 1, col=1)
            tables.append(Table(headers=headers, rows=entries, loc=loc))
    finally:
        wb.close()
    return Document(str(p), tables)


def _row_xlsx(values: list[str], p: Path, sheet: str, row_num: int, headers: list[str]) -> Row:
    cells: list[Cell] = []
    for j, value in enumerate(values):
        col = j + 1
        header = headers[j] if j < len(headers) else f"col{col}"
        cell = Cell(
            text=str(value).strip(),
            loc=SourceLoc(path=str(p), sheet=sheet, row=row_num, col=col, header=header),
        )
        cell.coerce()
        cells.append(cell)
    return Row(cells)


def document_from_records(records: list[dict[str, Any]], path: str | Path) -> Document:
    """Build a :class:`Document` from JSON records (e.g. an enterprise API reply).

    ``records`` is a list of flat dicts; the union of keys in first-seen order
    becomes the header row, and each record becomes one table row with source
    locations pointing at ``path``.
    """
    headers: list[str] = []
    seen: set[str] = set()
    for record in records:
        for key in record:
            if key not in seen:
                seen.add(key)
                headers.append(key)
    deduped = _dedupe_headers(headers)
    p = Path(str(path))
    rows = [
        _row_from_values([str(r.get(h, "")) for h in headers], p, i, deduped)
        for i, r in enumerate(records, start=2)
    ]
    table = Table(headers=deduped, rows=rows, loc=SourceLoc(path=str(p), row=2, col=1))
    return Document(str(p), [table])
