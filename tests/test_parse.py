import random
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from reconcheck.errors import TextDecodeError, UnsupportedFormatError
from reconcheck.parse import load_document


def test_csv_utf8(tmp_path: Path):
    p = tmp_path / "po.csv"
    p.write_text("料号,数量,金额\nA01,5 kg,100.00\n", encoding="utf-8")
    doc = load_document(p)
    table = doc.tables[0]
    assert table.headers == ["料号", "数量", "金额"]
    assert len(table.rows) == 1
    qty = table.rows[0].by_header("数量")
    assert qty.value == 5
    assert qty.unit == "kg"
    amount = table.rows[0].by_header("金额")
    assert amount.value == Decimal("100.00")
    loc = table.rows[0].by_header("料号").loc
    assert loc.row == 2 and loc.col == 1 and loc.header == "料号"
    assert loc.path == str(p)


def test_csv_gbk_fallback(tmp_path: Path):
    p = tmp_path / "gbk.csv"
    p.write_bytes("行号,名称\n1,螺丝\n".encode("gb18030"))
    doc = load_document(p)
    assert doc.tables[0].rows[0].by_header("名称").text == "螺丝"


def test_csv_ragged_row_gets_generated_headers(tmp_path: Path):
    p = tmp_path / "ragged.csv"
    p.write_text("a,b\n1,2,3\n", encoding="utf-8")
    table = load_document(p).tables[0]
    assert table.headers == ["a", "b", "col3"]
    assert table.rows[0].by_header("col3").text == "3"


def test_csv_empty_file(tmp_path: Path):
    p = tmp_path / "empty.csv"
    p.write_text("", encoding="utf-8")
    assert load_document(p).tables == []


def test_xlsx(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["料号", "数量"])
    ws.append(["A01", "3"])
    ws.append(["B02", "4"])
    x = tmp_path / "t.xlsx"
    wb.save(x)
    wb.close()
    doc = load_document(x)
    table = doc.tables[0]
    assert table.loc.sheet == "Sheet1"
    assert len(table.rows) == 2
    assert table.rows[0].by_header("料号").loc.row == 2
    assert table.rows[0].by_header("料号").loc.col == 1
    assert table.rows[1].by_header("数量").value == 4


def test_xlsx_sheet_filter(tmp_path: Path):
    wb = Workbook()
    a = wb.active
    a.title = "A"
    a.append(["k", "v"])
    a.append(["1", "10"])
    b = wb.create_sheet("B")
    b.append(["k", "v"])
    b.append(["1", "20"])
    x = tmp_path / "two.xlsx"
    wb.save(x)
    wb.close()
    doc = load_document(x, sheet="B")
    assert [t.loc.sheet for t in doc.tables] == ["B"]
    assert doc.tables[0].rows[0].by_header("v").value == 20


def test_unsupported_format(tmp_path: Path):
    p = tmp_path / "a.pdf"
    p.write_bytes(b"%PDF-1.4")
    with pytest.raises(UnsupportedFormatError):
        load_document(p)


def test_csv_binary_junk_raises_text_decode_error(tmp_path: Path):
    """Random bytes must fail the decode plausibility gate, not parse as junk."""
    p = tmp_path / "junk.csv"
    p.write_bytes(random.Random(42).randbytes(4096))
    with pytest.raises(TextDecodeError):
        load_document(p)


def test_csv_control_char_bytes_rejected(tmp_path: Path):
    # surrogates break utf-16/utf-8/gb18030; latin-1 fallback then fails the
    # control-char plausibility gate
    p = tmp_path / "ctl.csv"
    p.write_bytes(b"\x00\xd8\x00\x00" * 256)
    with pytest.raises(TextDecodeError):
        load_document(p)


def test_scientific_notation_parses(tmp_path: Path):
    p = tmp_path / "sci.csv"
    p.write_text("id,v\n1,1.5E3\n2,1e-2 kg\n", encoding="utf-8")
    table = load_document(p).tables[0]
    assert table.rows[0].by_header("v").value == Decimal("1500")
    small = table.rows[1].by_header("v")
    assert small.value == Decimal("0.01")
    assert small.unit == "kg"


def test_huge_number_to_dict_does_not_overflow(tmp_path: Path):
    p = tmp_path / "huge.csv"
    p.write_text(f"id,v\n1,{'9' * 400}\n", encoding="utf-8")
    cell = load_document(p).tables[0].rows[0].by_header("v")
    assert cell.value is not None  # Decimal holds arbitrary magnitude...
    assert cell.to_dict()["value"] is None  # ...but JSON float() must not crash


def test_dedupe_headers_never_shadows_real_suffix(tmp_path: Path):
    from reconcheck.parse.tabular import _dedupe_headers

    # a real "A_1" never shares a name with a generated one: all unique
    out = _dedupe_headers(["A", "A", "A_1"])
    assert out[0] == "A" and out[1] == "A_1"
    assert len(set(out)) == len(out) == 3
    # clean runs keep the friendly A / A_1 / A_2 numbering
    assert _dedupe_headers(["A", "A", "A"]) == ["A", "A_1", "A_2"]
    assert _dedupe_headers(["", "x", "x"]) == ["col1", "x", "x_1"]


def test_document_from_records_null_cells_become_empty():
    from reconcheck.parse.tabular import document_from_records

    doc = document_from_records(
        [{"id": "1", "note": None, "amt": "10"}],
        "records.json",
    )
    table = doc.tables[0]
    assert table.rows[0].by_header("note").text == ""
    assert table.rows[0].by_header("amt").text == "10"


def test_xlsx_corrupt_raises_engine_error_not_bare_zip(tmp_path: Path):
    p = tmp_path / "broken.xlsx"
    p.write_bytes(b"PK\x03\x04 not a real zip")
    with pytest.raises(UnsupportedFormatError):
        load_document(p)

def test_multichar_currency_prefixes_parse():
    from reconcheck.models import Cell, SourceLoc

    for text, want in [
        ("USD 12.30", Decimal("12.30")),
        ("-¥12.30", Decimal("-12.30")),
        ("HK$5", Decimal("5")),
        ("$ 1,200.00", Decimal("1200.00")),
        ("cost 12", None),  # a word that merely starts with a letter is no currency
    ]:
        c = Cell(text=text, loc=SourceLoc(path="x"))
        c.coerce()
        assert c.value == want, text


def test_xlsx_zip_bomb_guard(tmp_path, monkeypatch):
    import zipfile

    from reconcheck.parse import tabular

    p = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("xl/worksheets/sheet1.xml", "x" * 400)
    monkeypatch.setattr(tabular, "_MAX_XLSX_UNCOMPRESSED", 100)
    with pytest.raises(UnsupportedFormatError, match="uncompresses"):
        load_document(p)
